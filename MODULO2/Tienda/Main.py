from Tienda import Tienda
import random, math
from pathlib import Path
from openpyxl import load_workbook
import time

CARPETA = Path("C:/Users/luisd/Documents/Semestre 8/Topicos de IA/practicas/.venv/data")
ARCHIVO_DATOS = CARPETA / "datos_distribucion_tiendas.xlsx"
ARCHIVO_COSTOS = CARPETA / "matriz_costos_combustible.xlsx"  
ARCHIVO_DISTANCIAS = CARPETA / "matriz_distancias.xlsx" 
TIPO_CENTRO = "Centro de Distribución"

#datos
def leerNodos(ruta):
    
    hoja = load_workbook(ruta, read_only=True, data_only=True).active
    centros, tiendas = [], []
    for id, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=1):
        tipo, nombre, latitud, longitud = fila[0], fila[1], fila[2], fila[3]
        if tipo is None:
            continue
        nodo = Tienda(id, nombre, f"{longitud},{latitud}")   # mismo formato lon,lat de antes
        (centros if tipo == TIPO_CENTRO else tiendas).append(nodo)
    return centros, tiendas

def leerMatriz(ruta):
    hoja = load_workbook(ruta, read_only=True, data_only=True).active
    return [list(f) for f in hoja.iter_rows(min_row=2, values_only=True)]

#domino del problema

def temperaturaInicial(matriz, solucion, muestras=200):
  
    base = costo(matriz, solucion)
    deltas = [costo(matriz, generarVecinoAleatorio(solucion.copy())) - base for _ in range(muestras)]
    peores = [d for d in deltas if d > 0]
    
    return sum(peores) / len(peores) if peores else 1.0

def generarSolucionInicialSimple(centro, tiendas): #puntos a como fueron leídos
    return [centro] + tiendas



def swapRandom(arreglo):
    implicados = random.randint(2, len(arreglo) - 1)
 
    inicio = random.randint(1, len(arreglo) - implicados)   # nunca mueve la posición 0
    fin = inicio + implicados
 
    arreglo[inicio:fin] = arreglo[inicio:fin][::-1]
 
    return arreglo
def generarVecinoAleatorio(solucionActual):
    solucionCandidata = swapRandom(solucionActual)
    return solucionCandidata


def costo(matriz, solucion):
    total = 0.0
    for i in range(len(solucion) - 1):
        total += matriz[solucion[i].id - 1][solucion[i + 1].id - 1]
    total += matriz[solucion[-1].id - 1][solucion[0].id - 1]  # regreso al origen
    return total
      

def imprimirEdo(solucion):
    estado = " -> ".join(f"{tienda.id}: {tienda.nombre}" for tienda in solucion)
    print(f"Estado actual de la solución: {estado}")
    
    
#main
centros, tiendas = leerNodos(ARCHIVO_DATOS)
matrizCostos = leerMatriz(ARCHIVO_COSTOS)
matrizDistancias = leerMatriz(ARCHIVO_DISTANCIAS)

def recocidoSimulado(matriz, centro, tiendas, alpha=0.95, iteraciones=None):
    solucionActual = generarSolucionInicialSimple(centro, tiendas)
    costoActual = costo(matriz, solucionActual)
 
    T = temperaturaInicial(matriz, solucionActual)
    temperaturaFinal = T * 0.001
    if iteraciones is None:
        iteraciones = 1000                         # iteraciones por cada temperatura
 
    mejor, costoMejor = solucionActual[:], costoActual
    while T > temperaturaFinal:
        for _ in range(iteraciones):
            solucionCandidata = generarVecinoAleatorio(solucionActual.copy())
            costoCandidata = costo(matriz, solucionCandidata)
            delta = costoCandidata - costoActual
 
            if delta < 0 or random.random() < math.exp(-delta / T):
                solucionActual, costoActual = solucionCandidata, costoCandidata
                if costoActual < costoMejor:
                    mejor, costoMejor = solucionActual[:], costoActual
        T *= alpha
 
    return mejor, costoMejor


def imprimirRuta(solucion):
    ruta = " -> ".join(f"{t.id}: {t.nombre}" for t in solucion)
    print(f"{ruta} -> {solucion[0].id}: {solucion[0].nombre}")

print(f"{len(centros)} centros de distribución, {len(tiendas)} tiendas\n")
 
resultados = []
for centro in centros:
    inicio = time.time()
    mejor, costoMejor = recocidoSimulado(matrizCostos, centro, tiendas)
    km = costo(matrizDistancias, mejor)
    resultados.append((centro, mejor, costoMejor, km))
 
    print("\n")
    print(f"{centro.nombre}  |  costo combustible: {costoMejor:.2f}  |  {km:.2f} km  |  {time.time() - inicio:.1f} s")
    imprimirRuta(mejor)
 
print("\n")
print("RESUMEN")
for centro, mejor, costoMejor, km in sorted(resultados, key=lambda r: r[2]):
    print(f"{centro.nombre:<28} costo: {costoMejor:9.2f}   km: {km:9.2f}")